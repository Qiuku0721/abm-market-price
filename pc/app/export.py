"""CSV / Excel 导出。"""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def to_csv(records: list[dict]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(
        ["bullet_name", "price", "currency", "market_type",
         "captured_at", "source", "device_id", "snapshot_id"]
    )
    for r in records:
        writer.writerow([
            r.get("bullet_name", ""),
            r.get("price", ""),
            r.get("currency", "Koen"),
            r.get("market_type", "market"),
            r.get("captured_at", ""),
            r.get("source", ""),
            r.get("device_id", ""),
            r.get("snapshot_id", "") or "",
        ])
    return buf.getvalue()


def to_xlsx(records: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "价格记录"
    headers = ["子弹", "价格", "时间", "来源", "设备", "快照"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in records:
        ws.append([
            r.get("bullet_name", ""),
            r.get("price", ""),
            r.get("captured_at", ""),
            r.get("source", ""),
            r.get("device_id", ""),
            r.get("snapshot_id", "") or "",
        ])
    widths = [24, 12, 28, 12, 40, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
